"""Builds docs/slides.pptx from results/metrics.json, the figures and a live run of the demo.

    pip install python-pptx
    python docs/make_slides.py
"""
import json
import os
import subprocess
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
R = json.load(open(os.path.join(ROOT, "results/metrics.json")))
FIG = os.path.join(ROOT, "results/figures")
OUT = os.path.join(ROOT, "docs/slides.pptx")

NAVY, BLUE, INK, MUTED = RGBColor(0x14, 0x2B, 0x4F), RGBColor(0x2A, 0x78, 0xD6), RGBColor(0x1B, 0x1B, 0x1B), RGBColor(0x52, 0x51, 0x4E)
WHITE, LIGHT = RGBColor(255, 255, 255), RGBColor(0xEE, 0xF3, 0xFA)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def rect(s, x, y, w, h, fill):
    r = s.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    r.line.fill.background()
    r.fill.solid()
    r.fill.fore_color.rgb = fill
    return r


def text(s, x, y, w, h, lines, size=20, color=INK, bold=False, font="Calibri", bullet=False, space=8):
    tf = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines if isinstance(lines, list) else [lines]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = ("•  " + line) if bullet else line
        r.font.size, r.font.bold, r.font.name, r.font.color.rgb = Pt(size), bold, font, color
        p.space_after = Pt(space)


def slide(title, n):
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, 13.333, 1.05, NAVY)
    text(s, 0.55, 0.17, 12.2, 0.8, title, size=30, color=WHITE, bold=True)
    text(s, 12.4, 7.0, 0.8, 0.4, str(n), size=12, color=MUTED)
    return s


def picture(s, name, x, y, w=None, h=None):
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    s.shapes.add_picture(os.path.join(FIG, name), Inches(x), Inches(y), **kw)


def table(s, x, y, rows, widths, size=14, highlight_last=True):
    t = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y),
                           Inches(sum(widths)), Inches(0.4 * len(rows))).table
    for j, w in enumerate(widths):
        t.columns[j].width = Inches(w)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.text = val or " "
            p = c.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT
            r = p.runs[0]
            r.font.size = Pt(size)
            r.font.bold = i == 0 or (highlight_last and i == len(rows) - 1)
            r.font.color.rgb = WHITE if i == 0 else (BLUE if highlight_last and i == len(rows) - 1 else INK)
            c.fill.solid()
            c.fill.fore_color.rgb = NAVY if i == 0 else (LIGHT if i % 2 else WHITE)


def demo(*args):
    out = subprocess.run([sys.executable, os.path.join(ROOT, "src/predict.py"), *args],
                         capture_output=True, text=True, cwd=ROOT).stdout
    return [l for l in out.splitlines() if l.strip()]


e1, mean = R["E1_reproduce_paper"], R["E2_baselines"]["mean"]["test"]
sel = R["selected_model"]
wrong = R["E4_enriched_models"][f"{sel}|paper_alignment"]["test"]
diff_only, fin, cv = R["diff_only_model_test"], R["final_model_test"], R["cv5"]
e6, e7 = R["E6_proxy_validation"], R["E7_paired_cv_gradient_boosting"]["diff+size_before"]
c3, t = R["E3_feature_target_spearman"], R["target_summary_ms"]

# 1 title
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, 13.333, 7.5, NAVY)
text(s, 0.9, 2.0, 11.5, 1.8, "Predicting Bazel Build Times\nwith Machine Learning", size=46, color=WHITE, bold=True)
text(s, 0.9, 4.3, 11.5, 1.5, ["UE24CS352A Machine Learning, Mini-Project",
                              "Team: Deepak & Manaswin",
                              "Reproducing, auditing and improving Nadeem & Raza, Stanford CS229"],
     size=22, color=RGBColor(0xC9, 0xD6, 0xF0))

# 2 problem
s = slide("The problem", 2)
text(s, 0.7, 1.4, 12, 5.5, [
    "Bazel is incremental: a commit only re-runs the build actions whose inputs changed, so a build costs "
    "anything from under a second to tens of seconds.",
    "Task: predict the CPU time of the build that follows a commit, using only what is known before "
    "building, i.e. which files the commit changed.",
    "Use: CI schedulers can batch or prioritise builds; developers are warned that a change is expensive.",
    "The paper: linear regression on two categorical features (most common path prefix and file type, "
    "plus their cross) over 500 builds.",
    "Our goal: reproduce it, audit it, improve it.",
], size=22, bullet=True)

# 3 data
s = slide("Data: 500 real Bazel builds", 3)
text(s, 0.6, 1.35, 6.4, 5.5, [
    "500 consecutive commits of bazelbuild/bazel (2021); each built with bazel build //src:bazel-dev, CPU "
    "time read from the Build Event Protocol.",
    f"Target: median {t['50%'] / 1000:.1f} s, mean {t['mean'] / 1000:.1f} s, max {t['max'] / 1000:.1f} s; "
    "heavily right-skewed, so we model log(CPU time).",
    "Recovered the full list of changed files for all 500 commits from git, and actionsExecuted from the BEP files.",
    "Split as in the paper: builds 1-300 train, 301-400 validation, 401-500 test.",
], size=18, bullet=True)
picture(s, "target_distribution.png", 7.2, 1.5, w=5.7)

# 4 audit
s = slide("Auditing the original code", 4)
text(s, 0.7, 1.4, 12, 5.5, [
    f"Test bug: test_model() builds a new, untrained regressor, so every test prediction is 0 "
    f"(RMSE {e1['paper_test_bug_untrained_model_rmse']:,.0f} ms, the flat test line in the paper).",
    "Substring bugs: '.c' matches .cc and .css, '.h' matches .html, '.js' matches .json; a commit with no "
    "listed type defaults to JAVA; 'tools/' matches every com/google/devtools/ path.",
    f"Reproduced paper model: test ρ = {e1['with_cross']['test']['spearman']:.2f}, no better than predicting "
    f"the mean (RMSE {e1['with_cross']['test']['rmse_ms']:,.0f} vs {mean['rmse_ms']:,.0f} ms).",
    "Off-by-one: history was walked backwards with a warm cache, so build i re-does the work of the files "
    "changed by commit i-1, not by commit i.",
], size=21, bullet=True)

# 5 alignment
s = slide("The alignment fix: the signal appears", 5)
picture(s, "alignment_fix.png", 0.5, 1.35, w=7.6)
text(s, 8.4, 1.5, 4.6, 5.3, [
    "Same features, shifted by one commit.",
    f"touches src/main: ρ {c3['touches_src_main']['paper_alignment']:.2f} to "
    f"{c3['touches_src_main']['fixed_alignment']:.2f}.",
    f"only tests/docs: ρ {c3['only_tests_or_docs']['paper_alignment']:.2f} to "
    f"{c3['only_tests_or_docs']['fixed_alignment']:.2f}.",
    "Data provenance mattered more than the choice of model.",
], size=19, bullet=True)

# 6 features and models
s = slide("Features and models", 6)
text(s, 0.7, 1.4, 12, 5.5, [
    "Diff features (32): changed files per source area and per type (by real extension), log #files, "
    "log lines added/deleted, flags touches src/main, only tests/docs, touches BUILD/.bzl.",
    "Size features (8): total and largest bytes and imports/#includes of the changed files, overall and "
    "for Java and C/C++, read from git in the state before the commit (the version the build compiles).",
    "Models: ridge, random forest, gradient boosting (Huber loss) on log CPU time; first (cold-cache) build excluded.",
    "Feature set and model chosen together on validation RMSLE only; test scored once.",
    "100 test builds are few, so feature sets are also compared by paired 5-fold CV repeated 20 times.",
], size=20, bullet=True)

# 7 results
s = slide("Results on the test set (builds 401-500)", 7)
rows = [["Model", "RMSE ms", "MAE ms", "ρ", "AUC > 5 s"],
        ["Paper's test code (untrained)", f"{e1['paper_test_bug_untrained_model_rmse']:,.0f}", "", "", ""],
        ["Predict training mean", f"{mean['rmse_ms']:,.0f}", f"{mean['mae_ms']:,.0f}", "0.00", "0.50"]]
for name, m in [("Paper model reproduced", e1["with_cross"]["test"]), ("Diff features, paper alignment", wrong),
                ("Diff features, fixed alignment", diff_only), ("+ size features (final)", fin)]:
    rows.append([name, f"{m['rmse_ms']:,.0f}", f"{m['mae_ms']:,.0f}", f"{m['spearman']:.2f}", f"{m['auc_expensive']:.2f}"])
table(s, 0.6, 1.45, rows, [4.6, 1.4, 1.4, 1.0, 1.4], size=15)
text(s, 0.6, 5.0, 12.2, 2.0, [
    f"The final model cuts MAE by {100 * (1 - fin['mae_ms'] / e1['with_cross']['test']['mae_ms']):.0f}% against "
    f"the paper's model and flags builds over 5 s with AUC {fin['auc_expensive']:.2f}.",
    "Validation chose the size features; on these 100 test builds they are slightly behind diff-only, a gap "
    "well inside the fold-to-fold spread (next slide).",
], size=17, bullet=True)

# 8 cross-validation
s = slide("Cross-validation: the size features help consistently", 8)
table(s, 0.6, 1.45, [["5-fold CV, 499 builds", "RMSE ms", "ρ"],
                     ["Paper features (SGD)", f"{cv['paper_features_sgd']['rmse_ms'][0]:,.0f} ± {cv['paper_features_sgd']['rmse_ms'][1]:,.0f}",
                      f"{cv['paper_features_sgd']['spearman'][0]:.2f}"],
                     ["Diff features", "4,709 ± 834", "0.78"],
                     ["+ size features (final)", f"{cv['selected']['rmse_ms'][0]:,.0f} ± {cv['selected']['rmse_ms'][1]:,.0f}",
                      f"{cv['selected']['spearman'][0]:.2f}"]], [3.6, 2.0, 0.9], size=15)
picture(s, "size_features_paired.png", 7.0, 1.4, w=6.0)
text(s, 0.6, 4.0, 12.2, 3.0, [
    f"Paired on identical folds (20 x 5): RMSLE {e7['mean_diff']['rmsle']:+.3f}, RMSE "
    f"{e7['mean_diff']['rmse_ms']:+,.0f} ms, ρ {e7['mean_diff']['spearman']:+.3f}.",
    f"Better in {e7['folds_better']['rmsle']:.0%} of 100 folds and in "
    f"{e7['repetitions_better_rmsle']} of 20 repetitions.",
    "Before and after file state give the same gain.",
], size=18, bullet=True)

# 9 proxies
s = slide("Do hand-built cost estimates work?", 9)
picture(s, "proxy_validation.png", 0.5, 1.35, w=7.3)
text(s, 8.1, 1.45, 4.9, 5.4, [
    "Two proxies from a companion project that mined all 47k Bazel commits, where no real build time exists.",
    f"A: fixed language and directory weights, ρ {e6['proxy A: fixed formula']['spearman_cpu']:.2f}.",
    f"B: fitted on measured javac/g++ compile times, ρ {e6['proxy B: calibrated on javac/g++']['spearman_cpu']:.2f}.",
    f"Both lose to one learned feature, files in src/main/java (ρ {e6['files in src/main/java']['spearman_cpu']:.2f}), "
    "and do not help as model inputs.",
], size=17, bullet=True)

# 10 what drives cost
s = slide("What drives CPU time", 10)
picture(s, "cpu_vs_actions.png", 0.4, 1.35, w=6.3)
picture(s, "feature_importance.png", 6.9, 1.35, w=6.1)
text(s, 0.6, 5.7, 12.2, 1.5, [
    f"CPU time tracks the number of actions Bazel re-runs (ρ {R['bep_analysis']['spearman_cpu_vs_actions_executed']:.2f}); "
    f"{R['bep_analysis']['builds_with_1_action']} builds re-ran a single action.",
    "Which area a change touches matters most; how large the changed files are adds a smaller, consistent gain.",
], size=17, bullet=True)

# 11 demo
s = slide("Live demo: src/predict.py", 11)
lines = (["$ python src/predict.py --row 407"] + demo("--row", "407")[-2:]
         + ["", "$ python src/predict.py --row 415"] + demo("--row", "415")[-2:]
         + ["", "$ python src/predict.py --files site/docs/install.md"] + demo("--files", "site/docs/install.md")[-1:])
rect(s, 0.5, 1.35, 12.3, 4.2, RGBColor(0x1E, 0x1E, 0x2E))
text(s, 0.7, 1.45, 11.9, 4.0, lines, size=14, color=RGBColor(0xE6, 0xE6, 0xF0), font="Consolas", space=2)
text(s, 0.5, 5.75, 12.3, 1.5, [
    "Also works on a real commit: --repo <bazel clone> --commit <sha> (computes size features from git).",
    "With only a file list the sizes are unknown, so it says so and uses the diff-features model.",
], size=16, bullet=True)

# 12 conclusions
s = slide("Conclusions and limitations", 12)
text(s, 0.7, 1.4, 12, 5.6, [
    "The paper's model has no signal; three bugs and an off-by-one alignment explain why.",
    f"Fixed alignment + richer features: test ρ {fin['spearman']:.2f}, MAE "
    f"{100 * (1 - fin['mae_ms'] / e1['with_cross']['test']['mae_ms']):.0f}% lower, CV ρ {cv['selected']['spearman'][0]:.2f}.",
    "Size and dependency features of the changed files give a consistent gain in paired CV.",
    "Hand-built cost formulas are weaker than learned features.",
    "Limitations: 500 builds, one machine, one target; 100 test builds cannot separate models a few percent "
    "apart; the largest rebuilds are under-predicted.",
    "Next: reverse-dependency counts from bazel query 'rdeps(...)', and more builds.",
], size=20, bullet=True)

# 13 contributions
s = slide("Who did what", 13)
text(s, 0.7, 1.4, 5.8, 5.6, ["Deepak", ], size=24, bold=True, color=BLUE)
text(s, 0.7, 2.0, 5.8, 5.0, [
    "Recovered the 500 commit diffs and BEP metrics",
    "Audit: test bug, substring bugs, off-by-one alignment",
    "Diff features, SGD re-implementation, models",
    "Training pipeline, figures, demo CLI, first write-up",
], size=17, bullet=True)
text(s, 6.9, 1.4, 5.8, 5.6, ["Manaswin"], size=24, bold=True, color=BLUE)
text(s, 6.9, 2.0, 5.8, 5.0, [
    "Size and dependency features (scripts/size_features.py)",
    "Companion project: 47k-commit mining and two cost proxies, incl. measured compile-time calibration",
    "Proxy validation, paired repeated CV, joint feature-set selection",
    "Demo support for size features; updated README, write-up and slides",
], size=17, bullet=True)

prs.save(OUT)
print("wrote", OUT, f"({len(prs.slides)} slides)")
